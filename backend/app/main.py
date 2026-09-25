import hashlib
import io
import json
import unicodedata
import zipfile
import posixpath
from pathlib import Path
from typing import Literal
from fastapi import FastAPI, Depends, HTTPException, Request, Response, UploadFile, File, Form, Query
from fastapi.responses import FileResponse, JSONResponse
from sqlalchemy import select, or_, func, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session as DBSession,defer
from pypdf import PdfReader
from docx import Document
from .auth import admin, current_user, owner, hash_password, check_password, set_session
from .config import settings
from .db import get_db
from .models import User, Ticket, Material, Task, Result, ResultVersion, Audit, QueryRecord, PipelineJob, OutcomeSubmission, uid, now
from .schema import Credentials, SearchInput, FixedSearchInput, TaskInput, ResultEdit, RevisionInput, Evaluation,MaterialPurposeInput
from .query_prompt import PROMPT_VERSION, complete_prompt
from .evaluator import get_adapter
from .worker import search_text


app = FastAPI(title='成果影响力评价平台', version='0.1.0')
from .pipeline_api import router as pipeline_router
app.include_router(pipeline_router)
from .submission_api import router as submission_router
app.include_router(submission_router)


@app.get('/api/admin/pipeline-reports/{id}')
def pipeline_report(id: str, db: DBSession = Depends(get_db), user: User = Depends(admin)):
    from uuid import UUID
    from .pipeline_api import job_or_404,directory
    from .pipeline_store import PipelineStore,fingerprint
    from .skill_loader import contract
    from .codex_adapter import to_report
    try: validated_id=UUID(id)
    except ValueError: raise HTTPException(422,'工作流编号无效')
    job_or_404(db,validated_id)
    store=PipelineStore(directory(validated_id));state=store.read()
    stage=state['stages']['wu_evaluation']
    if stage['status']!='succeeded':raise HTTPException(409,'吴老师评价尚未完成，暂无可预览报告')
    path=(store.root/stage['output']).resolve()
    if not path.is_relative_to(store.root):raise HTTPException(409,'报告路径异常')
    value=json.loads(path.read_text(encoding='utf-8'))
    if fingerprint(value)!=stage['output_hash']:raise HTTPException(409,'报告内容与存档哈希不一致')
    assessment=contract.Assessment.model_validate(value['assessment'])
    report=to_report(assessment,[assessment.outcome_resolution.canonical_name or state['inputs']['title']])
    search_record=state['stages']['search_replay']
    if search_record['status']=='succeeded':
        search_path=(store.root/search_record['output']).resolve()
        if not search_path.is_relative_to(store.root):raise HTTPException(409,'检索底稿路径异常')
        search_output=json.loads(search_path.read_text(encoding='utf-8'))
        if fingerprint(search_output)!=search_record['output_hash']:raise HTTPException(409,'检索底稿与存档哈希不一致')
        raw=search_output.get('raw_result',{})
        if raw.get('review_cutoff'):report.review_window=f"外部证据核验截至 {raw['review_cutoff']}。{raw.get('cutoff_basis','')}"
    source_fingerprint=fingerprint({'wu_output':stage['output_hash'],'search_output':search_record.get('output_hash') if search_record['status']=='succeeded' else None})
    return {'pipeline_id':str(validated_id),'pipeline_status':state['status'],'review_status':'draft','payload':report.model_dump(),'source_fingerprint':source_fingerprint}


@app.post('/api/admin/pipeline-reports/{id}/draft',status_code=201)
def pipeline_report_draft(id: str,db: DBSession=Depends(get_db),user: User=Depends(admin)):
    from uuid import UUID
    from .pipeline_api import job_or_404
    try: identifier=UUID(id)
    except ValueError:raise HTTPException(422,'工作流编号无效')
    job_or_404(db,identifier,True)
    existing=db.scalar(select(Result).where(Result.pipeline_id==str(identifier)).with_for_update())
    response=pipeline_report(str(identifier),db,user)
    payload=response['payload']
    if existing:
        if existing.source_fingerprint==response['source_fingerprint']:return serialize(existing,('search_text',))
        if existing.source_fingerprint is None and Evaluation.model_validate(existing.payload).model_dump()==payload:
            existing.source_fingerprint=response['source_fingerprint'];db.commit();return serialize(existing,('search_text',))
        if existing.status!='draft':raise HTTPException(409,'工作流结果已更新，请先撤回已发布报告，再导入新草稿')
    if payload['evaluation_status'] in ('needs_scope_confirmation','insufficient_project_context'):
        raise HTTPException(409,'当前成果范围或上下文尚不明确，不能转入报告审核')
    if existing:
        row=existing;row.payload=payload;row.title=payload['title'];row.search_text=search_text(payload);row.revision+=1
        row.source_fingerprint=response['source_fingerprint']
    else:
        row=Result(pipeline_id=str(identifier),title=payload['title'],payload=payload,search_text=search_text(payload),source_fingerprint=response['source_fingerprint'])
        db.add(row);db.flush()
    db.add(ResultVersion(result_id=row.id,revision=row.revision,payload=payload,editor=user.id))
    audit(db,user,'pipeline_report_refresh' if existing else 'pipeline_report_draft',row.id);db.commit()
    return serialize(row,('search_text',))


@app.get('/api/admin/pipeline-reports/{id}/double-layer')
def pipeline_double_layer(id: str, db: DBSession = Depends(get_db), user: User = Depends(admin)):
    """Return the unified, reviewable double-layer evaluation package.

    The management layer and D1-D7 evidence layer share L1-L6 semantics while
    retaining independent evidence and provenance. They are deliberately not
    averaged or numerically merged.
    """
    from uuid import UUID
    from .pipeline_api import job_or_404, directory
    try:
        identifier = UUID(id)
    except ValueError:
        raise HTTPException(422, '工作流编号无效')
    job_or_404(db, identifier)
    state = __import__('app.pipeline_store', fromlist=['PipelineStore']).PipelineStore(directory(identifier)).read()
    stages = state['stages']
    if stages['wu_evaluation']['status'] != 'succeeded' or stages['v19_evaluation']['status'] != 'succeeded':
        raise HTTPException(409, '双层评价尚未完成，暂不可生成统一结果')
    import hashlib
    def read_stage(name):
        from .pipeline_store import fingerprint
        record = stages[name]
        path = (directory(identifier) / record['output']).resolve()
        if not path.is_relative_to(directory(identifier)):
            raise HTTPException(409, '阶段底稿路径异常')
        value = json.loads(path.read_text(encoding='utf-8'))
        if fingerprint(value) != record['output_hash']:
            raise HTTPException(409, '阶段底稿哈希校验失败')
        return value
    wu = read_stage('wu_evaluation')
    v19 = read_stage('v19_evaluation')
    assessment = wu.get('assessment') or {}
    result = v19.get('result') or {}
    run = result.get('project_synthesis') or result.get('project') or {}
    levels = {
        'management_level': (assessment.get('level_name') or assessment.get('level')),
        'dimension_level': (run.get('impact_level') or {}).get('level') if isinstance(run.get('impact_level'), dict) else None,
        'scope_level': (run.get('scope_impact_level') or {}).get('level') if isinstance(run.get('scope_impact_level'), dict) else None,
    }
    return {
        'evaluation_name': '双层影响力评价',
        'rubric_version': 'unified-double-layer-impact.v1',
        'pipeline_id': str(identifier), 'pipeline_status': state['status'],
        'review_status': 'draft', 'levels': levels,
        'management_layer': {'rubric_id': wu.get('rubric_id'), 'assessment': assessment},
        'dimension_layer': {'rubric_id': v19.get('rubric_id'), 'result': result},
        'rule': '两层共用L1-L6语义，证据独立保存；等级不相加、不平均，差异交管理员审核。',
    }


@app.middleware('http')
async def protect_origin(request: Request, call_next):
    if request.method not in ('GET', 'HEAD', 'OPTIONS'):
        origin = request.headers.get('origin')
        if origin and origin.rstrip('/') != settings().allowed_origin.rstrip('/'):
            return JSONResponse({'detail': '请求来源不受信任'}, status_code=403)
    response = await call_next(request)
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['Cache-Control'] = 'no-store'
    return response


def audit(db, user, action, target):
    db.add(Audit(actor=user.id, action=action, target=target))


def get_or_404(db, model, id):
    row = db.get(model, id)
    if row is None:
        raise HTTPException(404, '记录不存在')
    return row


def serialize(row, exclude=()):
    return {c.name: getattr(row, c.name) for c in row.__table__.columns if c.name not in exclude}


def normalize(query):
    return ' '.join(unicodedata.normalize('NFKC', query).casefold().split())


@app.get('/api/health')
def health(db: DBSession = Depends(get_db)):
    db.execute(text('SELECT 1'))
    return {'status': 'ok', 'adapter': settings().evaluation_adapter}


@app.get('/api/evaluation-schema')
def evaluation_schema():
    return Evaluation.model_json_schema()


@app.post('/api/auth/register', status_code=201)
def register(body: Credentials, response: Response, db: DBSession = Depends(get_db)):
    user = User(username=body.username, password_hash=hash_password(body.password))
    db.add(user)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, '账号已存在')
    set_session(response, user.id)
    return serialize(user, ('password_hash',))


@app.post('/api/auth/login')
def login(body: Credentials, response: Response, db: DBSession = Depends(get_db)):
    user = db.scalar(select(User).where(User.username == body.username))
    if not user or not check_password(body.password, user.password_hash):
        raise HTTPException(401, '账号或密码错误')
    set_session(response, user.id)
    return serialize(user, ('password_hash',))


@app.post('/api/auth/logout')
def logout(response: Response):
    response.delete_cookie('impact_session')
    return {'ok': True}


@app.get('/api/auth/me')
def me(user: User = Depends(current_user)):
    return serialize(user, ('password_hash',))


@app.get('/api/results')
def published(db: DBSession = Depends(get_db)):
    rows = db.scalars(select(Result).where(Result.status == 'published').order_by(Result.published_at.desc()).limit(30))
    return [serialize(x, ('search_text',)) for x in rows]


@app.get('/api/results/{id}')
def published_result(id: str, db: DBSession = Depends(get_db)):
    row = get_or_404(db, Result, id)
    if row.status != 'published':
        raise HTTPException(404, '记录不存在或尚未发布')
    return serialize(row, ('search_text',))


@app.get('/api/projects/suggestions')
def project_suggestions(q: str = Query(default='', max_length=200), db: DBSession = Depends(get_db)):
    statement = select(Result).where(Result.status == 'published')
    key = normalize(q)
    if key:
        statement = statement.where(or_(Result.title.contains(key, autoescape=True), Result.search_text.contains(key, autoescape=True)))
    rows = db.scalars(statement.order_by(Result.published_at.desc(), Result.id).limit(20))
    return [{'id': row.id, 'title': row.title, 'keywords': row.payload['keywords'],
             'category': row.payload['category'], 'is_demo': row.payload.get('is_demo', False)} for row in rows]


@app.post('/api/search')
def fixed_search(body: FixedSearchInput, request: Request, response: Response, db: DBSession = Depends(get_db)):
    row = db.scalar(select(Result).where(Result.id == str(body.project_id)).with_for_update())
    if not row or row.status != 'published':
        raise HTTPException(422, '所选成果不存在或已撤回，请从联想列表重新选择已发布成果')
    record = QueryRecord(owner=owner(request, response), result_id=row.id, result_revision=row.revision,
        project_title=row.title, task_type=body.task_type, detail=body.detail,
        prompt_version=PROMPT_VERSION, prompt=complete_prompt(row, body))
    db.add(record)
    db.commit()
    return {'results': [serialize(row, ('search_text',))], 'request_id': record.id,
            'task_type': record.task_type, 'detail': record.detail, 'mode': 'published_report',
            'message': '表单已保存，当前展示已有评价；未执行新的模型核验。'}


@app.get('/api/queries')
def my_queries(request: Request, response: Response, db: DBSession = Depends(get_db)):
    who = owner(request, response)
    rows = db.scalars(select(QueryRecord).where(QueryRecord.owner == who).order_by(QueryRecord.created_at.desc()).limit(100))
    return [serialize(row, ('owner', 'prompt')) for row in rows]


@app.get('/api/admin/queries')
def all_queries(db: DBSession = Depends(get_db), user: User = Depends(admin)):
    return [serialize(row, ('owner',)) for row in db.scalars(select(QueryRecord).order_by(QueryRecord.created_at.desc()).limit(200))]


@app.post('/api/admin/search')
def search(body: SearchInput, request: Request, response: Response, db: DBSession = Depends(get_db), user: User = Depends(admin)):
    query = normalize(body.query)
    if not query:
        raise HTTPException(422, '请输入有效关键词')
    # Contains uses escaped LIKE: user % and _ are literal characters.
    matches = list(db.scalars(select(Result).where(Result.status == 'published',
        or_(*[Result.search_text.contains(word, autoescape=True) for word in query.split()])).limit(50)))
    if matches:
        matches.sort(key=lambda x: (normalize(x.title) == query, sum(word in x.search_text for word in query.split())), reverse=True)
        return {'results': [serialize(x, ('search_text',)) for x in matches], 'ticket': None}
    who = owner(request, response)
    active_key = hashlib.sha256(f'{who}\0{query}'.encode()).hexdigest()
    ticket = db.scalar(select(Ticket).where(Ticket.active_key == active_key))
    if not ticket:
        ticket = Ticket(owner=who, query=body.query.strip(), normalized_query=query, active_key=active_key)
        db.add(ticket)
        try:
            db.commit()
        except IntegrityError:
            db.rollback()
            ticket = db.scalar(select(Ticket).where(Ticket.active_key == active_key))
    return {'results': [], 'ticket': serialize(ticket, ('owner', 'active_key'))}


@app.get('/api/tickets')
def my_tickets(request: Request, response: Response, db: DBSession = Depends(get_db)):
    who = owner(request, response)
    return [serialize(x, ('owner', 'active_key')) for x in db.scalars(select(Ticket).where(Ticket.owner == who).order_by(Ticket.created_at.desc()).limit(100))]


@app.get('/api/admin/overview')
def overview(db: DBSession = Depends(get_db), user: User = Depends(admin)):
    return {name: db.scalar(select(func.count()).select_from(model).where(condition)) for name, model, condition in [
        ('tickets', Ticket, Ticket.status.in_(['pending', 'processing'])),
        ('tasks', Task, Task.status.in_(['queued', 'running'])),
        ('drafts', Result, Result.status == 'draft'),
        ('published', Result, Result.status == 'published')]}


@app.get('/api/admin/tickets')
def all_tickets(db: DBSession = Depends(get_db), user: User = Depends(admin)):
    return [serialize(x, ('owner', 'active_key')) for x in db.scalars(select(Ticket).order_by(Ticket.created_at.desc()).limit(200))]


@app.post('/api/admin/tickets/{id}/close')
def close_ticket(id: str, db: DBSession = Depends(get_db), user: User = Depends(admin)):
    row = get_or_404(db, Ticket, id)
    if row.status != 'pending':
        raise HTTPException(409, '只能关闭尚未处理的工单')
    row.status, row.active_key = 'closed', None
    row.note = '管理员已关闭需求'
    audit(db, user, 'close_ticket', id)
    db.commit()
    return serialize(row, ('owner', 'active_key'))


@app.get('/api/admin/materials')
def materials(db: DBSession = Depends(get_db), user: User = Depends(admin)):
    return [serialize(x, ('text', 'storage_key')) for x in db.scalars(select(Material).options(defer(Material.text)).order_by(Material.created_at.desc()).limit(200))]


SUPPORTED_MATERIALS = ('.txt', '.md', '.pdf', '.docx', '.xlsx')
ARCHIVE_SUFFIXES = ('.zip',)


def _safe_archive_name(name: str):
    normalized = posixpath.normpath(name.replace('\\', '/')).lstrip('/')
    if not normalized or normalized == '.' or normalized.startswith('../') or '/..' in normalized or ':' in normalized:
        return None
    return normalized


def _read_upload(upload: UploadFile, limit: int):
    data = upload.file.read(limit + 1)
    if len(data) > limit:
        raise HTTPException(413, f'文件 {upload.filename or "未命名文件"} 超过 {limit // 1024 // 1024} MB 限制')
    return data


def _extract_text(filename: str, content: bytes):
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_MATERIALS:
        raise ValueError('不支持的类型；支持 TXT、Markdown、PDF、DOCX、XLSX')
    try:
        if suffix == '.pdf':
            reader = PdfReader(io.BytesIO(content))
            if len(reader.pages) > settings().pdf_max_pages:
                raise ValueError(f'PDF最多{settings().pdf_max_pages}页')
            page_texts=[page.extract_text() or '' for page in reader.pages]
            extracted = '\n'.join(f'[第{i+1}页]\n{text}' for i,text in enumerate(page_texts))
            has_text = any(text.strip() for text in page_texts)
        elif suffix == '.docx':
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError('解压后材料过大')
            doc = Document(io.BytesIO(content))
            extracted = '\n'.join([p.text for p in doc.paragraphs] + [' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows])
            has_text = bool(extracted.strip())
        elif suffix == '.xlsx':
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:raise ValueError('表格展开后超过100MB，请拆分')
            book=load_workbook(io.BytesIO(content),read_only=True,data_only=False,keep_links=False)
            lines=[];characters=0;cells_seen=0
            try:
                for sheet in book.worksheets:
                    if (sheet.max_row or 0)*(sheet.max_column or 0)>2000000:raise ValueError('工作表范围超过200万单元格，请清理无效空白格式或拆分')
                    for row in sheet.iter_rows():
                        cells_seen+=len(row)
                        if cells_seen>2000000:raise ValueError('工作簿超过200万单元格，请拆分')
                        for cell in row:
                            if cell.value is None:continue
                            value=str(cell.value)
                            if cell.data_type=='f':value='[公式，未执行或计算] '+value
                            line=f'[{sheet.title}!{cell.coordinate}] {value}'
                            if cell.data_type=='d':line+=f' [原显示格式：{cell.number_format}]'
                            characters+=len(line)+1
                            if characters>600000:raise ValueError('表格提取文本超过60万字，请拆分')
                            lines.append(line)
            finally:book.close()
            extracted='\n'.join(lines);has_text=bool(lines)
        else:
            extracted = content.decode('utf-8-sig')
            has_text = bool(extracted.strip())
        if not has_text:
            raise ValueError('未提取到文字，扫描件请先进行 OCR')
        if len(extracted) > 600000:
            raise ValueError('文本超过60万字，请拆分材料')
    except Exception as exc:
        raise HTTPException(422, f'材料解析失败：{str(exc)[:200]}')
    return extracted


def _store_material(db, user, filename: str, content: bytes, extracted: str):
    root = Path(settings().storage_dir).resolve()
    root.mkdir(parents=True, exist_ok=True)
    suffix = Path(filename).suffix.lower()
    key = uid() + suffix
    path = root / key
    try:
        path.write_bytes(content)
        row = Material(filename=filename, storage_key=key, sha256=hashlib.sha256(content).hexdigest(), size=len(content), text=extracted, uploaded_by=user.id)
        db.add(row)
        db.flush()
    except Exception:
        path.unlink(missing_ok=True)
        raise
    audit(db, user, 'upload_material', row.id)
    return row


def _material_error(filename, error):
    return {'filename': filename, 'status': 'rejected', 'error': str(error)}


@app.post('/api/admin/materials', status_code=201)
def upload(file: UploadFile = File(...), db: DBSession = Depends(get_db), user: User = Depends(admin)):
    filename = Path((file.filename or 'file').replace('\\', '/')).name[:255]
    if Path(filename).suffix.lower() in ARCHIVE_SUFFIXES:
        raise HTTPException(422, 'ZIP请使用批量上传入口；系统会解压后逐项处理')
    content = _read_upload(file, settings().max_upload_mb * 1024 * 1024)
    try:
        extracted = _extract_text(filename, content)
        row = _store_material(db, user, filename, content, extracted)
        db.commit()
    except HTTPException:
        db.rollback(); raise
    except Exception as exc:
        db.rollback(); raise HTTPException(422, f'{filename}解析失败：{str(exc)[:200]}')
    return serialize(row, ('storage_key', 'text'))


@app.post('/api/submission-materials/batch', status_code=207)
def upload_submission_materials(files: list[UploadFile] = File(...), db: DBSession = Depends(get_db), user: User = Depends(current_user)):
    return upload_batch(files, db, user, 'project')


@app.post('/api/admin/materials/batch', status_code=207)
def upload_batch(files: list[UploadFile] = File(...), db: DBSession = Depends(get_db), user: User = Depends(admin), purpose: Literal['project','reference','instruction'] = Form('project')):
    """Accept multiple files, directory uploads (webkitRelativePath) and ZIP archives.
    Each extracted supported file becomes a normal Material row; rejected entries do not abort valid siblings.
    """
    cfg = settings()
    if not files:
        raise HTTPException(422, '至少选择一个文件')
    if len(files) > cfg.batch_max_files:
        raise HTTPException(413, f'单次最多上传 {cfg.batch_max_files} 个文件')
    max_total = cfg.batch_max_total_mb * 1024 * 1024
    max_each = cfg.max_upload_mb * 1024 * 1024
    total_bytes = 0
    expanded_total = 0
    candidate_count = 0
    accepted = []
    rejected = []
    def accept_candidate(filename,content):
        nonlocal candidate_count
        if candidate_count>=cfg.batch_max_files:
            rejected.append(_material_error(filename,f'批次文件数超过 {cfg.batch_max_files} 个限制'))
            return
        candidate_count+=1
        try:
            extracted=_extract_text(filename,content)
            with db.begin_nested():
                row=_store_material(db,user,filename[:255],content,extracted)
                row.purpose=purpose
                db.flush()
            accepted.append(serialize(row,('storage_key','text'))|{'status':'accepted'})
        except Exception as exc:
            rejected.append(_material_error(filename,f'解析失败：{str(exc)[:200]}'))
    for upload_file in files:
        original = (upload_file.filename or 'file').replace('\\', '/')
        filename = original[:255]
        try:
            content = _read_upload(upload_file, max_each)
            total_bytes += len(content)
            if total_bytes > max_total:
                rejected.append(_material_error(filename, f'批次超过 {cfg.batch_max_total_mb} MB 总大小限制'))
                continue
            if Path(filename).suffix.lower() == '.zip':
                with zipfile.ZipFile(io.BytesIO(content)) as archive:
                    infos = [i for i in archive.infolist() if not i.is_dir()]
                    expanded = sum(i.file_size for i in infos)
                    if expanded_total + expanded > cfg.archive_max_expanded_mb * 1024 * 1024:
                        raise ValueError(f'本批压缩包展开总量超过 {cfg.archive_max_expanded_mb} MB 限制')
                    if candidate_count + len(infos) > cfg.batch_max_files:
                        raise ValueError(f'压缩包内文件数使批次超过 {cfg.batch_max_files} 个限制')
                    expanded_total+=expanded
                    for info in infos:
                        safe = _safe_archive_name(info.filename)
                        if not safe:
                            rejected.append(_material_error(f'{filename}!/{info.filename}', '压缩包路径不安全'))
                            continue
                        if Path(safe).suffix.lower() not in SUPPORTED_MATERIALS:
                            rejected.append(_material_error(f'{filename}!/{safe}', '压缩包内类型不支持'))
                            continue
                        if info.file_size > max_each:
                            rejected.append(_material_error(f'{filename}!/{safe}', '解压后单文件超过限制'))
                            continue
                        try:member=archive.read(info)
                        except Exception as exc:
                            rejected.append(_material_error(f'{filename}!/{safe}',f'压缩包条目读取失败：{str(exc)[:150]}'));continue
                        accept_candidate(f'{Path(filename).stem}/{safe}',member)
            elif Path(filename).suffix.lower() in SUPPORTED_MATERIALS:
                accept_candidate(filename,content)
            else:
                rejected.append(_material_error(filename, '不支持的类型；支持 TXT、Markdown、PDF、DOCX、XLSX 和 ZIP'))
        except zipfile.BadZipFile:
            rejected.append(_material_error(filename, 'ZIP文件损坏或不是有效压缩包'))
        except Exception as exc:
            rejected.append(_material_error(filename, exc))
    if accepted:
        db.commit()
    else:
        db.rollback()
    return {'accepted': accepted, 'rejected': rejected, 'accepted_count': len(accepted), 'rejected_count': len(rejected),
            'limits': {'max_files': cfg.batch_max_files, 'max_total_mb': cfg.batch_max_total_mb, 'max_each_mb': cfg.max_upload_mb, 'archive_expanded_mb': cfg.archive_max_expanded_mb}}


@app.patch('/api/admin/materials/{id}/purpose')
def classify_material(id: str, body: MaterialPurposeInput, db: DBSession = Depends(get_db), user: User = Depends(admin)):
    row=get_or_404(db,Material,id)
    row.purpose=body.purpose
    audit(db,user,'material_purpose_'+body.purpose,id);db.commit()
    return serialize(row,('storage_key','text'))


@app.get('/api/admin/materials/{id}')
def material_preview(id: str, db: DBSession = Depends(get_db), user: User = Depends(admin)):
    return serialize(get_or_404(db, Material, id), ('storage_key',))


@app.get('/api/admin/materials/{id}/download')
def download(id: str, db: DBSession = Depends(get_db), user: User = Depends(admin)):
    row = get_or_404(db, Material, id)
    path = Path(settings().storage_dir).resolve() / row.storage_key
    if not path.is_file():
        raise HTTPException(404, '原始文件未找到')
    return FileResponse(path, filename=row.filename, media_type='application/octet-stream')


@app.get('/api/admin/adapters')
def adapters(user: User = Depends(admin)):
    from .codex_adapter import CodexAdapter, codex_command
    reason=''
    try: codex_command()
    except ValueError as exc: reason=str(exc)
    return [
        {'name':'mock','version':'mock-v1','label':'模拟评价（仅用于流程演示）','available':True,'configured':settings().evaluation_adapter=='mock'},
        {'name':'codex','version':CodexAdapter().version,'label':'Codex · 吴老师v2成果评价','available':not reason,'reason':reason,'configured':settings().evaluation_adapter=='codex'},
    ]


@app.post('/api/admin/tasks', status_code=201)
def create_task(body: TaskInput, db: DBSession = Depends(get_db), user: User = Depends(admin)):
    if not body.title.strip() or any(not k.strip() or len(k) > 100 for k in body.keywords):
        raise HTTPException(422, '成果名称和关键词不能为空或过长')
    for id in body.material_ids:
        material=get_or_404(db, Material, id)
        if material.purpose!='project':raise HTTPException(422,'参考结果或流程说明不能作为本次评价的项目原始依据，请重新选择材料')
    if body.ticket_id:
        ticket = db.scalar(select(Ticket).where(Ticket.id == body.ticket_id).with_for_update())
        if not ticket:
            raise HTTPException(404, '工单不存在')
        if ticket.status != 'pending':
            raise HTTPException(409, '该工单已处理或正在处理中，请使用已有任务')
        ticket.status = 'processing'
    try:
        adapter = get_adapter(body.adapter or settings().evaluation_adapter)
        if adapter.name == 'codex':
            from .codex_adapter import codex_command
            codex_command()
    except ValueError as exc:
        raise HTTPException(503, str(exc))
    row = Task(**body.model_dump(exclude={'adapter'}), adapter=adapter.name, skill_version=adapter.version, model=adapter.model, created_by=user.id)
    db.add(row)
    db.flush()
    audit(db, user, 'create_task', row.id)
    db.commit()
    return serialize(row)


@app.get('/api/admin/tasks')
def tasks(db: DBSession = Depends(get_db), user: User = Depends(admin)):
    return [serialize(x, ('raw_output',)) for x in db.scalars(select(Task).order_by(Task.created_at.desc()).limit(200))]


@app.get('/api/admin/tasks/{id}/artifacts/{filename}')
def task_artifact(id: str, filename: str, db: DBSession = Depends(get_db), user: User = Depends(admin)):
    row=get_or_404(db,Task,id)
    root=Path(settings().storage_dir).resolve()/'evaluations'/row.id/f'attempt-{row.attempts}'
    artifacts={'evaluation-result.json','evidence-ledger.json','outcome-resolution.json','outcome-card.json','report.html','project-material-request.md','expert-review-task.md'}
    if filename in artifacts: path=root/'artifacts'/filename
    elif filename in {'stderr.log','events.jsonl','run.json','response.json'}: path=root/filename
    else: raise HTTPException(404,'不支持的底稿文件')
    if not path.is_file(): raise HTTPException(404,'此任务暂无该底稿文件')
    return FileResponse(path,filename=filename,media_type='application/octet-stream')


@app.post('/api/admin/tasks/{id}/retry')
def retry(id: str, db: DBSession = Depends(get_db), user: User = Depends(admin)):
    row = db.scalar(select(Task).where(Task.id == id).with_for_update())
    if not row:
        raise HTTPException(404, '任务不存在')
    if row.status != 'failed':
        raise HTTPException(409, '只有失败的任务可以重试')
    row.status, row.error, row.started_at, row.finished_at = 'queued', '', None, None
    audit(db, user, 'retry_task', id)
    db.commit()
    return serialize(row)


@app.get('/api/admin/results')
def results(db: DBSession = Depends(get_db), user: User = Depends(admin)):
    return [serialize(x, ('search_text',)) for x in db.scalars(select(Result).order_by(Result.created_at.desc()).limit(200))]


def locked_result(db, id, revision):
    row = db.scalar(select(Result).where(Result.id == id).with_for_update())
    if not row:
        raise HTTPException(404, '结果不存在')
    if row.revision != revision:
        raise HTTPException(409, '结果已被修改，请刷新后重试')
    return row


@app.put('/api/admin/results/{id}')
def edit_result(id: str, body: ResultEdit, db: DBSession = Depends(get_db), user: User = Depends(admin)):
    row = locked_result(db, id, body.revision)
    if row.status == 'published':
        raise HTTPException(409, '请先撤回已发布结果，再编辑')
    payload = body.payload.model_dump()
    if row.payload.get('is_demo') and not payload['is_demo']:
        raise HTTPException(422, '模拟结果不能移除演示标识')
    for ev in payload['evidence']:
        if ev['material_id']:
            get_or_404(db, Material, ev['material_id'])
    row.payload, row.title, row.search_text = payload, payload['title'], search_text(payload)
    row.revision += 1
    db.add(ResultVersion(result_id=row.id, revision=row.revision, payload=payload, editor=user.id))
    audit(db, user, 'edit_result', id)
    db.commit()
    return serialize(row, ('search_text',))


@app.post('/api/admin/results/{id}/publish')
def publish(id: str, body: RevisionInput, db: DBSession = Depends(get_db), user: User = Depends(admin)):
    preview=get_or_404(db,Result,id)
    if preview.pipeline_id:
        from .pipeline_api import job_or_404
        job_or_404(db,preview.pipeline_id,True)
    row = locked_result(db, id, body.revision)
    if row.pipeline_id:
        current=pipeline_report(row.pipeline_id,db,user)
        if row.source_fingerprint!=current['source_fingerprint']:
            raise HTTPException(409,'工作流来源已变化，请重新转入报告审核后发布，不能发布旧草稿')
    Evaluation.model_validate(row.payload)
    for evidence in row.payload.get('evidence',[]):
        if evidence.get('material_id'):
            material=get_or_404(db,Material,evidence['material_id'])
            if material.purpose!='project':raise HTTPException(409,'报告引用的材料已标记为参考或流程说明，请重新核验原始依据后发布')
    if row.payload.get('evaluation_status') in ('needs_scope_confirmation','insufficient_project_context'):
        raise HTTPException(409,'请先确认成果范围或补齐项目上下文并重新评价，当前草稿不可发布')
    original_task = db.get(Task,row.task_id) if row.task_id else None
    if original_task and original_task.adapter == 'codex' and original_task.raw_output:
        original_status=json.loads(original_task.raw_output).get('evaluation_status')
        if original_status in ('needs_scope_confirmation','insufficient_project_context'):
            raise HTTPException(409,'本次模型底稿要求补充上下文或确认范围，编辑展示JSON不能替代重新评价')
    row.status, row.published_at, row.reviewed_by = 'published', now(), user.id
    row.revision += 1
    db.add(ResultVersion(result_id=row.id, revision=row.revision, payload=row.payload, editor=user.id))
    # Exact aliases resolve matching requests from all users without exposing owners.
    aliases = {normalize(row.title), *[normalize(k) for k in row.payload['keywords']]}
    ticket_id = db.get(Task, row.task_id).ticket_id if row.task_id else None
    if row.pipeline_id:ticket_id=db.get(PipelineJob,row.pipeline_id).ticket_id
    for ticket in db.scalars(select(Ticket).where(Ticket.status.in_(['pending', 'processing', 'review']), or_((Ticket.status == 'pending') & Ticket.normalized_query.in_(aliases) & ~Ticket.id.in_(select(OutcomeSubmission.ticket_id)), Ticket.id == ticket_id))):
        ticket.status, ticket.result_id, ticket.active_key = 'published', row.id, None
    audit(db, user, 'publish_result', id)
    db.commit()
    return serialize(row, ('search_text',))


@app.post('/api/admin/results/{id}/withdraw')
def withdraw(id: str, body: RevisionInput, db: DBSession = Depends(get_db), user: User = Depends(admin)):
    row = locked_result(db, id, body.revision)
    row.status, row.published_at = 'draft', None
    row.revision += 1
    db.add(ResultVersion(result_id=row.id, revision=row.revision, payload=row.payload, editor=user.id))
    for ticket in db.scalars(select(Ticket).where(Ticket.result_id == id)):
        ticket.status, ticket.result_id, ticket.note = 'review', None, '结果已撤回，等待重新审核发布'
    audit(db, user, 'withdraw_result', id)
    db.commit()
    return serialize(row, ('search_text',))


@app.get('/api/admin/results/{id}/versions')
def versions(id: str, db: DBSession = Depends(get_db), user: User = Depends(admin)):
    get_or_404(db, Result, id)
    return [serialize(x) for x in db.scalars(select(ResultVersion).where(ResultVersion.result_id == id).order_by(ResultVersion.revision.desc()))]


@app.get('/api/admin/audit')
def audits(db: DBSession = Depends(get_db), user: User = Depends(admin)):
    return [serialize(x) for x in db.scalars(select(Audit).order_by(Audit.created_at.desc()).limit(100))]
