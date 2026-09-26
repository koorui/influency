"""Install the user's 09251930 grading document as self-contained skill resources."""
import json
from pathlib import Path
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT=Path(__file__).resolve().parents[1]
SOURCE=Path(r'C:\Users\32134\Desktop\科研成果六级影响力分级及理论依据-09251930.docx')
doc=Document(SOURCE)
parts=[]
for child in doc.element.body.iterchildren():
    if child.tag.endswith('}p'):
        text=Paragraph(child,doc).text.strip()
        if text:parts.append(text)
    elif child.tag.endswith('}tbl'):
        table=Table(child,doc)
        for n,row in enumerate(table.rows):
            parts.append('| '+' | '.join(c.text.replace('\n',' / ') for c in row.cells)+' |')
            if n==0:parts.append('| '+' | '.join('---' for c in row.cells)+' |')
        parts.append('')
levels={row.cells[0].text.split()[0]:{'name':row.cells[0].text.split(maxsplit=1)[1],
    'meaning':row.cells[1].text,'condition':row.cells[2].text} for row in doc.tables[0].rows[1:]}
dimensions={row.cells[0].text.split()[0]:{'name':row.cells[0].text.split(maxsplit=1)[1],
    'question':row.cells[1].text,'theory':row.cells[2].text} for row in doc.tables[1].rows[1:]}
policy='''## 本平台执行口径（用户于2026-09-25补充）

本文件是当前 L1–L6 和 D1–D7 的优先标准，覆盖旧版真实任务即L1、仅浦江接入算D6的限制。
对已经定位且具有原始材料的成果，每次完成评价必须给出 L1–L6 之一、七个 G1–G5 等级和明确理由，不以初步评价、待确认、待核验或空值替代等级。
按已提供证据能够支持的状态判级，不为提高等级补造事实。L1只要求成果形成与基本有效性验证；真实任务作用是L2、独立外部使用是L3、多主体持续使用是L4、稳定基础能力是L5、改变研究或工作方式是L6。
缺少高等级依据时选择较低等级。某维度没有可用外部证据时以G1作为本轮材料范围的保守基档，明确“本轮证据未建立更高等级”，不能写成现实中绝对不存在使用或认可。G1可以没有该维度引文，但必须解释证据范围；G2及以上必须有对应依据。空材料或成果身份不明属于工单输入问题，不能编造L1报告。
七维是整体L级的证据层，保留现有逐维G1–G5成熟度供底稿检查，不求和、不平均、不机械换算L级。该Word并未另给完整G1–G5逐档表，现有G档细则在不冲突处沿用。
D6考察稳定流程、平台或组织能力的集成，不限定浦江国家实验室；内部设计、实际输入输出、反馈闭环和持续运行分开。D7考察独立第三方评测、专业采用或权威认可；正式论文/专利是公开产出，不能单独证明外部认可。公开不等于复用，关注不等于认可。
判级以本轮提交材料可证明的成果状态为依据；项目启动、阶段截止、成果事件发生和材料形成日期分别记载。内部回顾材料可支持明确的技术事实，日期缺项写入适用边界，不清空已有等级，也不把截止日后或日期不明事件倒推为截止日前事实。公开来源的历史时间审计仍保留。
报告先写最终等级和关键事实，再说明适用范围和未达到更高等级的原因。低置信度可以保留，但不替代等级。人工验收只决定是否接受报告，不是生成明确等级的前置条件。
'''
standard={'version':'grading-20260925-1930','source':SOURCE.name,'levels':levels,'dimensions':dimensions,
          'completion_policy':'Every completed evaluation has explicit L1-L6 and seven G1-G5 grades; evidence gaps bound the conclusion, never fabricate facts.'}
for name in ['outcome-impact-evaluation','dual-layer-impact-v19']:
    refs=ROOT/'skills'/name/'references';refs.mkdir(exist_ok=True)
    (refs/'grading-standard-20260925.md').write_text('# 当前判级标准\n\n'+policy+'\n## 用户提供的原文\n\n'+'\n\n'.join(parts)+'\n',encoding='utf-8')
    (refs/'grading-standard-20260925.json').write_text(json.dumps(standard,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print('Installed L1-L6 and D1-D7 standard in both evaluation skills.')
